# -*- coding: utf-8 -*-
"""작품소개자료(.docx) 생성. 실행: python report.py → submission/ 에 저장, 이어서 python topdf.py 로 PDF 변환"""
import pathlib
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = pathlib.Path(__file__).resolve().parent.parent
CAP = ROOT / "captures"
OUT = ROOT / "submission"; OUT.mkdir(exist_ok=True)
NAME = "[경희대학교] Team Wispy_김현정_작품소개자료"
URL = "https://wispy-khu.vercel.app"
FONT = "맑은 고딕"
PURPLE = RGBColor(0x7C, 0x5C, 0xD6)
BLACK = RGBColor(0, 0, 0); GREY = RGBColor(0x59, 0x59, 0x59)
BORDER = "7F7F7F"; HEAD = "F2F2F2"
BASE = 8.8
W = 182

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(14)
sec.top_margin, sec.bottom_margin = Mm(11), Mm(9)
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(BASE)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.13
_ppr = st.element.get_or_add_pPr()
for _k in ("autoSpaceDE", "autoSpaceDN"):
    _e = OxmlElement("w:" + _k); _e.set(qn("w:val"), "0"); _ppr.append(_e)

C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
MID = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def run(p, text, bold=False, color=BLACK, size=None):
    r = p.add_run(text); r.bold = bold; r.font.color.rgb = color
    r.font.name = FONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size: r.font.size = Pt(size)
    return r


def para(text="", size=None, before=0, after=0, align=None, container=None, indent=0):
    p = (container or doc).add_paragraph()
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent: pf.left_indent = Mm(indent)
    if align is not None: p.alignment = align
    if text: run(p, text, size=size)
    return p


def item(text, container=None, after=1.0):
    """본문 항목: ' - 문장' 형태"""
    p = (container or doc).add_paragraph()
    pf = p.paragraph_format; pf.left_indent = Mm(5.2); pf.first_line_indent = Mm(-2.6); pf.space_after = Pt(after)
    run(p, "- " + text)
    return p


def h1(text, newpage=False):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.page_break_before = newpage
    pf.space_before = Pt(7); pf.space_after = Pt(3); pf.keep_with_next = True
    run(p, text, bold=True, size=11.5)
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")): b.set(qn("w:" + k), v)
    bd.append(b); pPr.append(bd)


def h2(text, before=4, container=None, first=False):
    c = container or doc
    p = c.paragraphs[0] if (first and container is not None) else c.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True
    run(p, "□ " + text, bold=True, size=9.6)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)


def borders(table, on=True, color=BORDER, sz="4", only=None):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); show = on and (only is None or edge in only)
        e.set(qn("w:val"), "single" if show else "nil"); e.set(qn("w:sz"), sz); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def cellpad(table, top=26, bottom=26, left=70, right=70):
    tblPr = table._tbl.tblPr; m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def mktable(rows, cols, widths, container=None):
    t = (container or doc).add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for r in t.rows:
        for i, c in enumerate(r.cells): c.width = Mm(widths[i])
    return t


def grid(data, widths, size=8.2, center_first=True):
    """일반 표: 전체 테두리, 회색 머리글"""
    t = mktable(len(data), len(widths), widths); borders(t); cellpad(t)
    for ri, row in enumerate(data):
        for ci, txt in enumerate(row):
            c = t.cell(ri, ci); c.vertical_alignment = MID
            p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.08
            if ri == 0:
                shade(c, HEAD); p.alignment = C; run(p, txt, bold=True, size=size)
            elif ci == 0 and center_first:
                p.alignment = C; run(p, txt, size=size)
            else:
                run(p, txt, size=size)
    return t


def gap(pt=3):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = Pt(pt)


# ───────────── 제목 ─────────────
t = mktable(1, 1, [W]); borders(t, color="000000", sz="12", only=("top", "bottom")); cellpad(t, 60, 70, 0, 0)
c = t.cell(0, 0)
p = c.paragraphs[0]; p.alignment = C
run(p, "Wispy(위스피)", bold=True, size=19)
p = c.add_paragraph(); p.alignment = C; p.paragraph_format.space_before = Pt(1)
run(p, "스마트폰에 빠져 계속 보고 있을 때 먼저 전화를 걸어오는 AI 캐릭터 서비스", size=10)
p = c.add_paragraph(); p.alignment = C; p.paragraph_format.space_before = Pt(3)
run(p, "경희대학교 Team Wispy", size=9); run(p, "   |   ", size=9, color=GREY)
run(p, "막지 않습니다. 다시 선택하게 합니다.", bold=True, size=9, color=PURPLE)

# ───────────── ① ─────────────
h1("① 팀 소개, 작품 소개")
t = mktable(1, 2, [144, 38]); borders(t, on=False); cellpad(t, 0, 0, 0, 30)
c = t.cell(0, 0)
h2("팀 소개", before=0, container=c, first=True)
item("Team Wispy는 경희대학교 학생 5명으로 구성된 팀입니다. 2026년 8월 한 달 동안 기획, 디자인, 개발, 배포를 직접 진행했습니다.", container=c)
h2("작품 소개", container=c)
item("Wispy는 스마트폰 사용을 막는 앱이 아닙니다. 잠깐만 보려던 스마트폰에 빠져 계속 보고 있으면 AI 캐릭터 ‘위스피’가 사용자에게 전화를 겁니다.", container=c)
item("전화를 받으면 위스피가 “지금 뭐 보고 있었어요?”라고 말을 걸고, 30~50초 정도 대화한 뒤 “이제 뭐 할 거예요?”를 한 번 묻습니다. 더 볼지 그만 볼지는 사용자가 정합니다.", container=c)
item("화면을 보고 있는 사람에게는 화면 속 경고가 잘 전달되지 않습니다. 그래서 시각 대신 청각, 즉 벨소리와 목소리로 스크롤을 한 번 멈추게 하는 것이 이 작품의 출발점입니다.", container=c)
item("처음에 입력한 관심사와 ‘요즘 하려는 일’, 지난 통화 내용을 대화에 반영하며, 훈계하거나 그만 보라고 말하지 않습니다.", container=c)
item("모바일 웹앱(PWA)으로 만들어 설치와 회원가입 없이 QR로 접속해 바로 사용할 수 있습니다.", container=c)
c = t.cell(0, 1); c.vertical_alignment = MID
p = c.paragraphs[0]; p.alignment = C
p.add_run().add_picture(str(CAP / "07_demo_incoming_call_banner.png"), width=Mm(29))
p = c.add_paragraph(); p.alignment = C; run(p, "[그림 1] 전화 수신 화면", size=7.6)

h2("사용 흐름")
t = mktable(1, 1, [W]); borders(t); cellpad(t, 50, 50, 70, 70)
p = t.cell(0, 0).paragraphs[0]; p.alignment = C
run(p, "개인화 입력  →  사용 감지  →  위스피 알림  →  AI 전화  →  짧은 대화  →  다시 선택  →  기록·리포트", size=8.8)

# ───────────── ② ─────────────
h1("② 작품 개요")
h2("개발 배경", before=0)
item("과학기술정보통신부의 「2025년 스마트폰 과의존 실태조사」에 따르면 스마트폰 이용자의 22.7%가 과의존 위험군이며, 20~30대는 29.5%로 전체보다 높습니다.")
item("스크린타임 통계와 앱 잠금 기능은 이미 널리 쓰이고 있습니다. 그러나 경고가 떠도 ‘1분 더 보기’를 누르거나 제한을 해제하면 되기 때문에, 이미 몰입한 사용자의 행동은 잘 바뀌지 않습니다.")
item("저희는 문제를 ‘사용시간이 길다’가 아니라 ‘처음 의도와 실제 행동이 다르다’로 정의했습니다. 보려고 고른 영화 두 시간과 잠깐 쉬려다 지나간 숏폼 두 시간은 스크린타임이 같아도 같은 문제가 아니라고 보았습니다.")
gap(2)
grid([
    ["구분", "기존 스크린타임·차단 앱", "Wispy"],
    ["문제 정의", "사용시간이 깁니다", "의도와 실제 행동이 어긋납니다"],
    ["개입 방법", "통계를 보여주거나 시간이 되면 앱을 잠급니다", "캐릭터가 전화를 걸어 짧게 대화합니다"],
    ["사용 감각", "시각 (화면 위 경고)", "청각 (벨소리, 목소리)"],
    ["결정 주체", "앱이 사용을 막습니다", "사용자가 직접 선택합니다"],
], [26, 78, 78])
h2("목표 시스템")
item("대상은 스마트폰 사용을 줄이고 싶지만 막상 그 순간에는 조절이 어려운 20~34세 대학생과 직장인입니다.")
item("사용자가 정한 시간이 지나면 먼저 알림으로 말을 걸고, 이어서 전화로 개입합니다.")
item("통화는 30~50초 안에 끝나는 것을 목표로 하며, 다음 행동을 묻는 질문은 통화당 한 번만 합니다.")
item("개인화 입력, 알림, 전화, 대화, 기록·리포트가 하나의 흐름으로 이어지도록 구성합니다.")

# ───────────── ③ ─────────────
h1("③ 작품 구성 및 상세 내용", newpage=True)
grid([
    ["단계", "기능", "내용"],
    ["개인화 입력", "온보딩 8단계", "되찾고 싶은 시간대, 대상 앱(6종)과 제한 시간, 관심사, ‘요즘 하려는 일’을 입력합니다."],
    ["사용 감지", "이탈 시간 확인", "앱을 벗어난 시각을 저장하고, 제한 시간을 넘겨 돌아오면 전화 화면을 띄웁니다."],
    ["위스피 알림", "알림 문구 생성", "관심사와 할 일을 반영한 한 줄 문구를 매번 새로 생성합니다."],
    ["AI 전화", "실시간 음성 통화", "발신자가 ‘위스피’로 표시되며, 말하는 도중 끼어들기와 음소거, 스피커 전환을 지원합니다."],
    ["짧은 대화", "4단계 대화", "“지금 뭐 보고 있었어요?” → 대화 → “이제 뭐 할 거예요?”(1회) → 마무리 인사 순입니다."],
    ["다시 선택", "통화 종료", "인사가 끝나면 위스피가 먼저 전화를 끊고, 스트레칭 등 다음 행동을 제안합니다."],
    ["기록·리포트", "통화 기록, 주간 리포트", "통화별 앱·시간·마지막 대화를 기록하고, 주간 횟수와 요일별·앱별 비중을 보여줍니다."],
    ["Memory", "이전 통화 기억", "통화 끝부분의 대화를 저장해 다음 통화의 프롬프트에 반영합니다."],
], [24, 34, 124], size=8.1)
h2("시스템 구조")
p = para(align=C); p.add_run().add_picture(str(CAP / "architecture.png"), width=Mm(140))
p = para(align=C, after=1); run(p, "[그림 2] 시스템 아키텍처", size=7.6)

# ───────────── ④ ─────────────
h1("④ 개발 세부 내용")
h2("사용 기술", before=0)
item("React 19, Vite 8, Tailwind CSS 4 / Express 4, TypeScript, Vercel Serverless / OpenAI Realtime API(WebRTC), gpt-4o-mini / Upstash Redis")
h2("개발 중 문제와 해결")
grid([
    ["문제", "해결 방법"],
    ["API 키를 브라우저에 둘 수 없음", "서버리스 함수가 임시 세션 토큰만 발급받아 전달하고, 브라우저는 이 토큰으로 OpenAI와 WebRTC로 직접 연결합니다. 키는 서버에만 있고 음성은 서버를 거치지 않습니다."],
    ["웹에서 다른 앱 사용시간 확인 불가", "visibilitychange 이벤트가 발생하면 앱을 벗어난 시각을 localStorage에 저장하고, 돌아왔을 때 경과 시간을 계산합니다. 백그라운드에서 코드가 멈춰도 정확히 계산됩니다."],
    ["백그라운드 알림이 오지 않음", "브라우저가 백그라운드 탭을 멈추면 알림을 보장할 수 없어, 체험 모드와 홈 화면의 벨 버튼으로 같은 흐름을 바로 확인할 수 있게 했습니다."],
    ["통화가 몰리면 429 오류 발생", "분당 토큰 한도가 따로 적용되는 두 모델(gpt-realtime, gpt-realtime-mini)에 통화를 번갈아 배정하고, 부하 테스트로 동시 통화를 직접 측정했습니다."],
    ["전화 연결이 느림", "홈 화면에 있는 동안 세션을 미리 발급받고, 연결 시 토큰 발급과 마이크 권한 요청을 동시에 진행합니다."],
    ["첫 인사가 중간에 끊김", "AI 목소리가 마이크로 들어가 끼어들기로 처리된 것이 원인이어서, 첫 인사가 끝날 때까지 마이크 입력을 차단했습니다."],
    ["긴 통화에서 말투가 흐트러짐", "매 턴마다 말투 지시와 사용자가 입력한 계획을 다시 전달하도록 했습니다."],
    ["인사 도중에 전화가 끊김", "AI가 종료 함수(end_call)를 인사 음성과 함께 호출하는 경우가 있어, 재생이 끝난 뒤 종료하게 했습니다."],
    ["로그인 없이 개인화 필요", "브라우저별 익명 ID로 Redis에 프로필을 저장하고, 허용 도메인·공유 시크릿·호출 제한을 적용했습니다."],
], [44, 138], size=8.0, center_first=False)

# ───────────── ⑤ ─────────────
h1("⑤ 구현 결과", newpage=True)
shots = [("03_onboarding_personalize.png", "개인화 입력"), ("04_home.png", "홈"),
         ("06_demo_notification_banner.png", "체험 모드 알림"), ("09_call_active.png", "통화 중"),
         ("10_log.png", "통화 기록"), ("11_report.png", "주간 리포트")]
t = mktable(2, 6, [W / 6] * 6); borders(t, on=False); cellpad(t, 10, 10, 10, 10)
for i, (f, cap) in enumerate(shots):
    c = t.cell(0, i); p = c.paragraphs[0]; p.alignment = C
    p.add_run().add_picture(str(CAP / f), width=Mm(28.5))
    c = t.cell(1, i); p = c.paragraphs[0]; p.alignment = C
    run(p, f"({i+1}) {cap}", size=8)
p = para(align=C, before=1, after=3); run(p, "[그림 3] 주요 화면", size=7.6)

t = mktable(1, 2, [32, 150]); borders(t, on=False); cellpad(t, 0, 0, 0, 40)
c = t.cell(0, 0); c.vertical_alignment = MID
p = c.paragraphs[0]; p.alignment = C; p.add_run().add_picture(str(CAP / "qr_wispy.png"), width=Mm(25))
p = c.add_paragraph(); p.alignment = C; run(p, "wispy-khu.vercel.app", size=7.4)
c = t.cell(0, 1); c.vertical_alignment = MID
h2("배포 및 체험 방법", before=0, container=c, first=True)
item(f"{URL} 에 배포했습니다. 휴대폰으로 QR을 찍으면 바로 접속됩니다.", container=c)
item("관심사와 할 일을 입력한 뒤 홈에서 ‘진짜처럼 체험해보기’를 누르면 알림, 전화 수신, 통화, 기록까지 한 번에 확인할 수 있습니다.", container=c)
item("2026년 8월에 개발했으며 커밋은 141건입니다. 주요 화면 9개와 API 7개로 구성되어 있습니다.", container=c)
item("프롬프트, 음성 톤, 통화 길이, 동시 접속을 확인하는 테스트 스크립트 7개를 만들어 반복해서 개선했습니다.", container=c)
h2("구현 현황", before=5)
grid([
    ["구분", "현재 구현된 기능", "향후 계획"],
    ["감지", "앱을 벗어난 시간으로 감지하고, 돌아오면 전화 화면을 표시합니다.", "앱별 실제 사용시간을 감지하고 사용 패턴에 따라 자동으로 개입합니다."],
    ["알림", "개인화 알림 문구를 생성하고 로컬 알림으로 표시합니다.", "구현해 둔 서버 Web Push 발송 기능을 감지 흐름에 연결합니다."],
    ["통화", "WebRTC 실시간 음성, 끼어들기, AI 주도 종료를 지원합니다.", "사용자마다 개입 시점과 빈도를 조정합니다."],
    ["개인화", "관심사와 할 일, 이전 통화 내용을 반영합니다.", "LLM으로 통화를 요약하고 더 긴 기간을 기억합니다."],
    ["기록", "통화 기록과 주간 리포트를 제공합니다.", "통화 후 실제 행동을 측정하고 리포트 공유 기능을 추가합니다."],
], [18, 82, 82], size=8.1)

# ───────────── ⑥ ─────────────
h1("⑥ 기대 효과")
item("앱이 막아서 멈추는 것과 통화 후 스스로 정해서 멈추는 것은 사용자에게 남는 경험이 다릅니다. 잠금을 해제하는 식의 우회가 줄고, 스스로 조절했다는 경험이 쌓일 것으로 기대합니다.", after=1.5)
item("설치와 가입이 필요 없어 학교 행사나 과의존 예방 교육 현장에서 QR만으로 체험할 수 있습니다.", after=1.5)
item("음성은 브라우저와 OpenAI가 직접 주고받고 서버는 토큰 발급과 프로필 저장만 담당합니다. 서버 부담이 작아 소규모 팀도 운영할 수 있는 구조입니다.", after=1.5)
item("수익 모델은 Freemium입니다. 목표 설정, 알림, 기록·리포트는 무료로 제공하고, 개인화 AI 전화와 Memory, 상세 리포트는 Premium(월 2,900원)으로 제공할 계획입니다.", after=1.5)
h2("향후 추진 계획", before=5)
grid([
    ["0~3개월", "3~6개월", "6~12개월"],
    ["현재 MVP로 전화 개입의 수용도, 통화 후 행동 변화, 재사용 의향을 검증합니다.", "앱별 실제 사용시간 감지를 적용하고 개입·반응 데이터를 축적합니다.", "사용자별 개입 시점과 방식을 최적화하고 Premium 모델을 검증합니다."],
], [60.6, 60.7, 60.7], size=8.1, center_first=False)

docx_path = OUT / (NAME + ".docx")
doc.save(str(docx_path)); print("saved", docx_path)
